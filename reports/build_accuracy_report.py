from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "reports" / "BIS_Accuracy_Robustness_Report.docx"
BASELINE = json.loads((ROOT / "reports" / "evaluation" / "robustness_summary.json").read_text(encoding="utf-8"))
AFTER = json.loads((ROOT / "reports" / "evaluation_after_aliases" / "robustness_summary.json").read_text(encoding="utf-8"))


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(9)


def add_table(document: Document, headers: list[str], rows: list[list[object]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Light Shading Accent 1"
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], str(value))
    document.add_paragraph()


def pct(value: object) -> str:
    return f"{float(value):.2f}%"


def main() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)
    styles["Title"].font.name = "Aptos Display"
    styles["Title"].font.size = Pt(24)
    styles["Heading 1"].font.color.rgb = RGBColor(31, 78, 121)
    styles["Heading 2"].font.color.rgb = RGBColor(47, 84, 150)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("BIS Standards Recommendation Engine\nAccuracy & Robustness Report")
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(31, 78, 121)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Prepared from local evaluation runs on 2 May 2026").italic = True
    doc.add_paragraph()

    doc.add_heading("Executive Summary", level=1)
    doc.add_paragraph(
        "The system is very strong on the official public validation shape and on English/Hinglish product descriptions. "
        "The original robustness run found a real gap in Spanish/French translation and sparse, vague product wording. "
        "A deterministic multilingual alias-expansion layer was added to the query processor and rerun against the same 200-query suite."
    )
    add_table(
        doc,
        ["Evaluation", "Queries", "Hit Rate @3", "MRR @5", "Avg Latency", "Status"],
        [
            ["Official public set", 10, "100.00%", "1.0000", "0.04-0.06s", "Schema OK"],
            [
                "Robustness baseline",
                BASELINE["total_queries"],
                pct(BASELINE["hit_rate_at_3"]),
                f"{BASELINE['mrr_at_5']:.4f}",
                f"{BASELINE['avg_latency_seconds']:.4f}s",
                "Before aliases",
            ],
            [
                "Robustness after aliases",
                AFTER["total_queries"],
                pct(AFTER["hit_rate_at_3"]),
                f"{AFTER['mrr_at_5']:.4f}",
                f"{AFTER['avg_latency_seconds']:.4f}s",
                "After aliases",
            ],
        ],
    )

    doc.add_heading("Scope And Method", level=1)
    doc.add_paragraph(
        "The test suite generated 200 labeled queries from the 10 public target standards. "
        "Each target received 20 query variants: public-style English, short telegraphic input, typo/noisy input, direct IS-code mention, Hinglish, Hindi Devanagari, Spanish, French, Tamil transliteration, and vague low-context wording."
    )
    doc.add_paragraph(
        "Metrics match the hackathon automated scoring emphasis: Hit Rate @3, MRR @5, and latency. "
        "The official public test was also run through inference.py and eval_script.py to verify strict judge compatibility."
    )
    doc.add_paragraph(
        "Reference context came from the provided rulebook DOCX and a web check of the public hackathon page, which confirms the Building Materials focus, top 3-5 standard output requirement, and scoring on Hit Rate @3, MRR @5, latency, technical excellence, innovation, usability, and presentation."
    )

    doc.add_heading("System Under Test", level=1)
    doc.add_paragraph(
        "The current pipeline parses BIS SP 21 material into structured standards, builds BM25 and FAISS artifacts, normalizes and expands queries, then ranks candidates with deterministic metadata and product-family boosts. "
        "Dense FAISS artifacts are present but dense retrieval is disabled by default in the current HybridRetriever configuration, so the measured path is primarily sparse retrieval plus deterministic ranking logic."
    )
    doc.add_paragraph(
        "The newly added query expansion layer maps common multilingual product phrases and noisy aliases back to the English product terms already represented in the catalog. "
        "This keeps inference deterministic and fast while directly addressing hidden-set risks from multilingual MSE-style wording."
    )

    doc.add_heading("Robustness By Query Style", level=1)
    style_rows = []
    for style, metrics in AFTER["by_style"].items():
        before = BASELINE["by_style"].get(style, {})
        style_rows.append(
            [
                style,
                metrics["queries"],
                pct(before.get("hit_rate_at_3", 0)),
                pct(metrics["hit_rate_at_3"]),
                f"{before.get('mrr_at_5', 0):.4f}",
                f"{metrics['mrr_at_5']:.4f}",
            ]
        )
    add_table(doc, ["Style", "Queries", "Baseline HR@3", "After HR@3", "Baseline MRR", "After MRR"], style_rows)

    doc.add_heading("Robustness By Target Standard", level=1)
    target_rows = []
    for target, metrics in AFTER["by_target"].items():
        before = BASELINE["by_target"].get(target, {})
        target_rows.append(
            [
                target,
                metrics["queries"],
                pct(before.get("hit_rate_at_3", 0)),
                pct(metrics["hit_rate_at_3"]),
                f"{before.get('mrr_at_5', 0):.4f}",
                f"{metrics['mrr_at_5']:.4f}",
            ]
        )
    add_table(doc, ["Target", "Queries", "Baseline HR@3", "After HR@3", "Baseline MRR", "After MRR"], target_rows)

    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_heading("What The Results Mean", level=1)
    doc.add_paragraph(
        "Strengths: the engine is highly accurate for judge-like English product descriptions, short fragments, Hinglish, direct-code queries, and the known cement/concrete/aggregate categories. "
        "Latency is far below the 5-second target, with the post-fix 200-query average at 0.0296 seconds and p95 at 0.0561 seconds."
    )
    doc.add_paragraph(
        "Original weaknesses: before the alias layer, Spanish and French queries scored only 20% Hit Rate @3, and underspecified queries scored 55%. "
        "Those failures came from vocabulary mismatch rather than slow ranking or schema issues."
    )
    doc.add_paragraph(
        "Post-fix status: the same 200-query suite reached 100% Hit Rate @3 and 0.9942 MRR@5. "
        "Only two queries were not rank-1; both still placed the expected standard within top 3."
    )

    doc.add_heading("Residual Risks", level=1)
    risks = [
        "The 200-query robustness suite is synthetic and label-controlled; it is valuable for stress testing but does not replace the hidden private set.",
        "The multilingual layer covers likely product phrases for the public target families, not every Indian language or every possible spelling.",
        "Dense retrieval is currently disabled by default, so unrelated hidden products with low lexical overlap may still need broader semantic recall.",
        "The report evaluates retrieval codes; if a demo adds natural-language rationales, those rationales should be checked for hallucinated standards.",
    ]
    for risk in risks:
        doc.add_paragraph(risk, style="List Bullet")

    doc.add_heading("Recommendations Before Submission", level=1)
    recommendations = [
        "Keep inference.py, eval_script.py, requirements.txt, and README command examples exactly judge-compatible.",
        "Expand the multilingual alias table with more Indian-language transliterations for cement, steel, concrete, aggregate, pipe, roofing, masonry, and block terms.",
        "Add a small no-hallucination guard in the demo/rationale path so every mentioned IS code must come from retrieved_standards.",
        "Document the unique parser/chunking strategy clearly in the deck: PDF parsing to structured catalog, chunk_text enrichment, BM25/FAISS artifacts, RRF-ready design, deterministic boosts, and multilingual aliases.",
        "Consider a final private-style smoke set with unseen building-material standards beyond the 10 public examples to reduce overfitting risk.",
    ]
    for item in recommendations:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("Evidence Files", level=1)
    evidence_rows = [
        ["reports/evaluation/robustness_results.json", "Baseline 200-query output"],
        ["reports/evaluation/robustness_results.csv", "Baseline query-level CSV"],
        ["reports/evaluation_after_aliases/robustness_results.json", "Post-fix 200-query output"],
        ["reports/evaluation_after_aliases/robustness_results.csv", "Post-fix query-level CSV"],
        ["reports/evaluation_after_aliases/public_results.json", "Official public-set judge-format output"],
        ["scripts/robustness_evaluation.py", "Reproducible evaluation harness"],
    ]
    add_table(doc, ["Path", "Purpose"], evidence_rows)

    doc.add_heading("External References Checked", level=1)
    refs = [
        "Unstop hackathon page: https://unstop.com/hackathons/bureau-of-indian-standards-x-sigma-squad-ai-hackathon-indian-institute-of-technology-iit-tirupati-1679162/amp",
        "Faiss repository: https://github.com/facebookresearch/faiss",
        "BEIR benchmark paper: https://arxiv.org/abs/2104.08663",
    ]
    for ref in refs:
        doc.add_paragraph(ref, style="List Bullet")

    doc.add_paragraph()
    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Conclusion: ready for public metrics; continue broadening multilingual coverage for hidden-set resilience.").bold = True

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
